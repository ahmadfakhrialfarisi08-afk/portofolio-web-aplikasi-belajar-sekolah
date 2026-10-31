"""
DOCX Renderer: Document Schema (schema.validasi_dokumen) -> file .docx.

Tidak tahu apa-apa soal AI, Flask, atau penyimpanan file. Input: dict schema bersih + identitas
(yang tersedia saja) + template gaya. Output: bytes.

Dependency: python-docx (import dilakukan saat render, supaya kalau belum terpasang, bagian lain
aplikasi tetap jalan dan route mengembalikan pesan jelas lewat RendererTidakTersedia).
"""
import io
import re
from datetime import datetime

from .schema import SchemaError, bersihkan_teks, validasi_dokumen
from .templates import BULAN, ambil_template, gabung_template


class RendererTidakTersedia(RuntimeError):
    """python-docx belum terpasang di server."""


def _muat_docx():
    try:
        import docx  # noqa: F401
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as e:  # pragma: no cover - tergantung lingkungan
        raise RendererTidakTersedia("Library python-docx belum terpasang (pip install python-docx).") from e
    return {
        "Document": docx.Document, "Cm": Cm, "Pt": Pt, "RGBColor": RGBColor, "qn": qn,
        "OxmlElement": OxmlElement, "ALIGN": WD_ALIGN_PARAGRAPH, "TAB": WD_TAB_ALIGNMENT,
        "TABLE_ALIGN": WD_TABLE_ALIGNMENT,
    }


# ---------------------------------------------------------------------------
# Helper XML kecil
# ---------------------------------------------------------------------------
def _set_font(style_atau_run, nama, X):
    """Set font + buang atribut tema (asciiTheme dst.) yang diam-diam menimpa nama font."""
    style_atau_run.font.name = nama
    rpr = style_atau_run.element.get_or_add_rPr() if hasattr(style_atau_run, "element") else style_atau_run._r.get_or_add_rPr()
    rfonts = rpr.find(X["qn"]("w:rFonts"))
    if rfonts is None:
        rfonts = X["OxmlElement"]("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        k = X["qn"](a)
        if k in rfonts.attrib:
            del rfonts.attrib[k]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(X["qn"](a), nama)


# Elemen yang HARUS muncul setelah w:pBdr / w:tblBorders / w:shd menurut skema OOXML.
# Menempelkan elemen di akhir (append) menghasilkan urutan tidak valid; Word bisa menolak file.
_SETELAH_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                 "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                 "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                 "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_SETELAH_TBLBORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
                       "w:tblPrChange")
_SETELAH_TCSHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                  "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def _garis_paragraf(p, sisi, warna, X, tebal, jarak):
    """Garis pada satu sisi paragraf ('top'/'bottom'), disisipkan pada urutan XML yang benar."""
    ppr = p._p.get_or_add_pPr()
    bdr = X["OxmlElement"]("w:pBdr")
    b = X["OxmlElement"]("w:" + sisi)
    b.set(X["qn"]("w:val"), "single")
    b.set(X["qn"]("w:sz"), str(tebal))
    b.set(X["qn"]("w:space"), str(jarak))
    b.set(X["qn"]("w:color"), warna)
    bdr.append(b)
    ppr.insert_element_before(bdr, *_SETELAH_PBDR)


def _garis_bawah_paragraf(p, warna, X, tebal=6):
    _garis_paragraf(p, "bottom", warna, X, tebal, 1)


def _garis_atas_paragraf(p, warna, X):
    _garis_paragraf(p, "top", warna, X, 4, 4)


def _tambah_field(par, kode, ukuran, X):
    """Field Word (PAGE / NUMPAGES) sebagai rangkaian run begin/instr/separate/hasil/end.
    Teks cadangan '1' tampil sampai Word menghitung ulang."""
    def _run():
        r = par.add_run()
        r.font.size = ukuran
        return r

    r = _run()
    fc = X["OxmlElement"]("w:fldChar")
    fc.set(X["qn"]("w:fldCharType"), "begin")
    r._r.append(fc)

    r = _run()
    it = X["OxmlElement"]("w:instrText")
    it.set(X["qn"]("xml:space"), "preserve")
    it.text = " %s " % kode
    r._r.append(it)

    r = _run()
    fc = X["OxmlElement"]("w:fldChar")
    fc.set(X["qn"]("w:fldCharType"), "separate")
    r._r.append(fc)

    r = _run()
    r.text = "1"

    r = _run()
    fc = X["OxmlElement"]("w:fldChar")
    fc.set(X["qn"]("w:fldCharType"), "end")
    r._r.append(fc)


def _num_id_baru(doc, nama_style, X):
    """Buat instance penomoran baru (mulai dari 1) untuk gaya 'List Number'.
    Tanpa ini, semua daftar bernomor di dokumen akan melanjutkan nomor daftar sebelumnya."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    ppr = doc.styles[nama_style].element.pPr
    lama = ppr.numPr.numId.val
    abstrak = None
    for num in numbering.num_lst:
        if num.numId == lama:
            abstrak = num.abstractNumId.val
            break
    if abstrak is None:
        return None
    baru = numbering.add_num(abstrak)
    baru.add_lvlOverride(ilvl=0).add_startOverride(1)
    return baru.numId


def _pasang_num(p, num_id, X):
    ppr = p._p.get_or_add_pPr()
    numpr = ppr.get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = num_id


_RE_HEX6 = re.compile(r"^[0-9A-Fa-f]{6}$")


def _warna(hex_):
    return hex_.upper() if isinstance(hex_, str) and _RE_HEX6.match(hex_) else "000000"


def _bayangi_sel(sel, fill, X):
    tcpr = sel._tc.get_or_add_tcPr()
    shd = X["OxmlElement"]("w:shd")
    shd.set(X["qn"]("w:val"), "clear")
    shd.set(X["qn"]("w:color"), "auto")
    shd.set(X["qn"]("w:fill"), fill)
    tcpr.insert_element_before(shd, *_SETELAH_TCSHD)


def _tanpa_border_tabel(tabel, X):
    tblpr = tabel._tbl.tblPr
    borders = X["OxmlElement"]("w:tblBorders")
    for sisi in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = X["OxmlElement"]("w:" + sisi)
        e.set(X["qn"]("w:val"), "nil")
        borders.append(e)
    tblpr.insert_element_before(borders, *_SETELAH_TBLBORDERS)


def _tabel_tetap(tabel, lebar_cm, X):
    """Layout tetap + lebar kolom konsisten di tblGrid DAN tiap sel (Word & LibreOffice sama-sama patuh)."""
    tabel.autofit = False  # python-docx menaruh w:tblLayout pada urutan yang benar
    for i, w in enumerate(lebar_cm):
        tabel.columns[i].width = X["Cm"](w)
        for sel in tabel.columns[i].cells:
            sel.width = X["Cm"](w)


def _lebar_kolom(kolom, baris, total_cm):
    """Lebar kolom proporsional isi (dengan batas bawah), jumlah = total_cm."""
    n = len(kolom)
    bobot = []
    for i in range(n):
        panjang = [len(kolom[i])] + [len(r[i]) for r in baris]
        rata = sum(panjang) / len(panjang)
        bobot.append(min(max(rata, 4.0), 55.0) ** 0.8)
    jumlah = sum(bobot)
    minimal = 1.4
    lebar = [max(minimal, total_cm * b / jumlah) for b in bobot]
    # normalisasi supaya total tepat total_cm
    faktor = total_cm / sum(lebar)
    return [w * faktor for w in lebar]


def _tanggal_indonesia(dt=None):
    dt = dt or datetime.now()
    return "%d %s %d" % (dt.day, BULAN[dt.month - 1], dt.year)


# ---------------------------------------------------------------------------
# Render utama
# ---------------------------------------------------------------------------
def _bersihkan_identitas(identitas):
    """Hanya nilai str yang tidak kosong, dibersihkan dari karakter kontrol (python-docx menolaknya)."""
    hasil = {}
    for k, v in (identitas or {}).items():
        if isinstance(v, str):
            t = bersihkan_teks(v, 200).replace("\n", " ").strip()
            if t:
                hasil[str(k)] = t
    return hasil


_GARIS_ISIAN = "." * 40
MAKS_META_TAMPIL = 20


def _gabung_meta(dok, identitas):
    """Meta dari schema + identitas yang tersedia + isian_kosong (baris bertitik untuk diisi tangan).
    Data yang tidak tersedia TIDAK dikarang: hanya diberi garis titik."""
    meta = [list(m) for m in dok.get("meta", [])]
    kunci = lambda k: k.strip().lower()
    ada = {kunci(k) for k, _ in meta}

    def tambah(nama, nilai):
        if kunci(nama) not in ada:
            meta.append([nama, nilai])
            ada.add(kunci(nama))

    if identitas.get("guru_nama") and not any(k.startswith(("penyusun", "guru", "nama guru")) for k in ada):
        tambah("Penyusun", identitas["guru_nama"])
    if identitas.get("mapel") and not any(k.startswith(("mata pelajaran", "mapel")) for k in ada):
        tambah("Mata Pelajaran", identitas["mapel"])

    # Pemetaan nama isian -> kunci identitas; kalau datanya tersedia, tidak perlu dikosongkan.
    peta = (("nip", "nip"), ("nama guru", "guru_nama"), ("guru", "guru_nama"), ("penyusun", "guru_nama"),
            ("sekolah", "sekolah"), ("mata pelajaran", "mapel"), ("mapel", "mapel"))
    for nama in dok.get("isian_kosong", []):
        n = kunci(nama)
        if any(n.startswith(awal) and identitas.get(kid) for awal, kid in peta):
            continue
        tambah(nama, _GARIS_ISIAN)
    return meta[:MAKS_META_TAMPIL]



def render_docx(dok, identitas=None, template=None, tanggal=None):
    """dok: dict hasil schema.validasi_dokumen. identitas: {"sekolah","guru_nama","nip","mapel"...}
    (hanya yang tersedia). template: dict gaya parsial (digabung di atas template bawaan jenis dokumen).
    tanggal: datetime untuk tanda tangan (bawaan: sekarang). Return bytes .docx."""
    X = _muat_docx()
    identitas = _bersihkan_identitas(identitas)
    tpl = ambil_template(dok.get("jenis"))
    if template:
        tpl = gabung_template(tpl, template)
    Cm, Pt, RGB, ALIGN = X["Cm"], X["Pt"], X["RGBColor"], X["ALIGN"]

    doc = X["Document"]()
    hal = tpl["halaman"]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(hal["lebar"]), Cm(hal["tinggi"])
    sec.left_margin, sec.right_margin = Cm(hal["kiri"]), Cm(hal["kanan"])
    sec.top_margin, sec.bottom_margin = Cm(hal["atas"]), Cm(hal["bawah"])
    sec.header_distance = sec.footer_distance = Cm(1.1)
    lebar_teks = hal["lebar"] - hal["kiri"] - hal["kanan"]

    # --- Gaya ---
    normal = doc.styles["Normal"]
    _set_font(normal, tpl["font"], X)
    normal.font.size = Pt(tpl["ukuran_isi"])
    normal.paragraph_format.space_after = Pt(tpl["spasi_setelah"])
    normal.paragraph_format.line_spacing = tpl["spasi_baris"]

    h1 = doc.styles["Heading 1"]
    _set_font(h1, tpl["font"], X)
    h1.font.size, h1.font.bold = Pt(tpl["ukuran_h1"]), True
    h1.font.color.rgb = RGB.from_string(_warna(tpl["warna_h1"]))
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(14), Pt(4)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    _set_font(h2, tpl["font"], X)
    h2.font.size, h2.font.bold, h2.font.italic = Pt(tpl["ukuran_h2"]), True, False
    h2.font.color.rgb = RGB(0, 0, 0)
    h2.paragraph_format.space_before, h2.paragraph_format.space_after = Pt(8), Pt(3)
    h2.paragraph_format.keep_with_next = True

    for nama in ("List Bullet", "List Number"):
        _set_font(doc.styles[nama], tpl["font"], X)
        doc.styles[nama].paragraph_format.space_after = Pt(3)

    # --- Header & footer ---
    if tpl.get("header_footer"):
        hp = sec.header.paragraphs[0]
        kiri = identitas.get("sekolah", "")
        kanan = dok["judul"] if len(dok["judul"]) <= 55 else dok["judul"][:52].rstrip() + "..."
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(lebar_teks), X["TAB"].RIGHT)
        r = hp.add_run(kiri + "\t\t" + kanan)  # 2 tab: lewati tab tengah bawaan style Header, mendarat di tab kanan
        _set_font(r, tpl["font"], X)
        r.font.size = Pt(9)
        r.font.color.rgb = RGB(0x59, 0x59, 0x59)
        _garis_bawah_paragraf(hp, _warna(tpl["warna_garis"]), X)

        fp = sec.footer.paragraphs[0]
        fp.alignment = ALIGN.CENTER
        _garis_atas_paragraf(fp, _warna(tpl["warna_garis"]), X)
        for teks, kode in (("Halaman ", "PAGE"), (" dari ", "NUMPAGES")):
            r1 = fp.add_run(teks)
            r1.font.size = Pt(9)
            _tambah_field(fp, kode, Pt(9), X)

    # --- Judul dokumen ---
    judul = dok["judul"].upper() if tpl.get("judul_kapital") else dok["judul"]
    p = doc.add_paragraph()
    p.alignment = ALIGN.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(judul)
    r.bold = True
    r.font.size = Pt(tpl["ukuran_judul"])
    r.font.color.rgb = RGB.from_string(_warna(tpl["warna_h1"]))
    if identitas.get("sekolah"):
        p2 = doc.add_paragraph()
        p2.alignment = ALIGN.CENTER
        p2.paragraph_format.space_after = Pt(10)
        r = p2.add_run(identitas["sekolah"])
        r.font.size = Pt(tpl["ukuran_isi"] + 1)
        _garis_bawah_paragraf(p2, _warna(tpl["warna_garis"]), X, tebal=8)

    # --- Identitas / meta (hanya yang tersedia) ---
    meta = _gabung_meta(dok, identitas)
    if meta:
        if tpl.get("meta_sebagai_tabel"):
            t = doc.add_table(rows=0, cols=2)
            _tanpa_border_tabel(t, X)
            lk = min(5.0, lebar_teks * 0.3)
            for k, v in meta:
                cells = t.add_row().cells
                for c, teks, tebal in ((cells[0], k, True), (cells[1], ": " + v, False)):
                    par = c.paragraphs[0]
                    par.paragraph_format.space_after = Pt(1)
                    rr = par.add_run(teks)
                    rr.bold = tebal
                    rr.font.size = Pt(tpl["ukuran_tabel"] + 0.5)
            _tabel_tetap(t, (lk, lebar_teks - lk), X)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        else:
            for k, v in meta:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(1)
                par.add_run(k + ": ").bold = True
                par.add_run(v)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- Isi ---
    for bg in dok["bagian"]:
        if bg["judul"]:
            doc.add_heading(bg["judul"], level=1)
        for b in bg["blok"]:
            tipe = b["tipe"]
            if tipe == "paragraf":
                for baris in b["teks"].split("\n"):
                    if baris.strip():
                        par = doc.add_paragraph(baris.strip())
                        par.alignment = ALIGN.JUSTIFY
            elif tipe == "subjudul":
                doc.add_heading(b["teks"], level=2)
            elif tipe == "daftar":
                if b["gaya"] == "nomor":
                    num_id = _num_id_baru(doc, "List Number", X)
                    for it in b["item"]:
                        par = doc.add_paragraph(it, style="List Number")
                        if num_id is not None:
                            _pasang_num(par, num_id, X)
                else:
                    for it in b["item"]:
                        doc.add_paragraph(it, style="List Bullet")
            elif tipe == "tabel":
                _render_tabel(doc, b, tpl, lebar_teks, X)

    # --- Tanda tangan ---
    if dok.get("tanda_tangan"):
        _render_tanda_tangan(doc, identitas, tpl, lebar_teks, tanggal, X)

    # --- Properti file ---
    cp = doc.core_properties
    cp.title = dok["judul"][:200]
    cp.author = identitas.get("guru_nama") or "SCORUS AI"
    cp.subject = dok.get("jenis", "")
    cp.comments = ""
    cp.language = "id-ID"
    cp.last_modified_by = cp.author
    cp.created = cp.modified = tanggal or datetime.now()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_tabel(doc, b, tpl, lebar_teks, X):
    Cm, Pt, ALIGN = X["Cm"], X["Pt"], X["ALIGN"]
    kolom, baris = b["kolom"], b["baris"]
    lebar = _lebar_kolom(kolom, baris, lebar_teks)
    t = doc.add_table(rows=1, cols=len(kolom))
    t.style = "Table Grid"
    t.alignment = X["TABLE_ALIGN"].CENTER
    # baris judul: diulang di tiap halaman
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = X["OxmlElement"]("w:tblHeader")
    trpr.append(th)
    for i, nama in enumerate(kolom):
        c = t.rows[0].cells[i]
        c.width = Cm(lebar[i])
        _bayangi_sel(c, _warna(tpl["warna_kepala_tabel"]), X)
        par = c.paragraphs[0]
        par.alignment = ALIGN.CENTER
        par.paragraph_format.space_after = Pt(2)
        rr = par.add_run(nama)
        rr.bold = True
        rr.font.size = Pt(tpl["ukuran_tabel"])
    for r in baris:
        row = t.add_row()
        cs = X["OxmlElement"]("w:cantSplit")
        row._tr.get_or_add_trPr().append(cs)
        for i, teks in enumerate(r):
            c = row.cells[i]
            c.width = Cm(lebar[i])
            potong = teks.split("\n") if teks else [""]
            for j, bagian in enumerate(potong):
                par = c.paragraphs[0] if j == 0 else c.add_paragraph()
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.line_spacing = 1.05
                if len(teks) <= 4 and teks:  # nomor / skor pendek: rata tengah
                    par.alignment = ALIGN.CENTER
                rr = par.add_run(bagian)
                rr.font.size = Pt(tpl["ukuran_tabel"])
    _tabel_tetap(t, lebar, X)
    # jarak setelah tabel
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


def _render_tanda_tangan(doc, identitas, tpl, lebar_teks, tanggal, X):
    Cm, Pt, ALIGN = X["Cm"], X["Pt"], X["ALIGN"]
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    t = doc.add_table(rows=1, cols=2)
    _tanpa_border_tabel(t, X)
    _tabel_tetap(t, (lebar_teks / 2, lebar_teks / 2), X)
    t.rows[0]._tr.get_or_add_trPr().append(X["OxmlElement"]("w:cantSplit"))  # blok TTD tidak terbelah halaman
    kiri, kanan = t.rows[0].cells
    # Nama yang tidak tersedia dibiarkan garis kosong (tidak dikarang).
    garis = "(" + " " * 34 + ")"
    blok = (
        (kiri, ["Mengetahui,", "Kepala Sekolah"], garis),
        (kanan, [_tanggal_indonesia(tanggal), "Guru Mata Pelajaran"], identitas.get("guru_nama") or garis),
    )
    for sel, baris_atas, nama in blok:
        first = True
        for teks in baris_atas:
            par = sel.paragraphs[0] if first else sel.add_paragraph()
            first = False
            par.alignment = ALIGN.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.add_run(teks)
        for _ in range(3):
            sel.add_paragraph().paragraph_format.space_after = Pt(0)
        par = sel.add_paragraph()
        par.alignment = ALIGN.CENTER
        rr = par.add_run(nama)
        rr.bold = nama != garis
        if identitas.get("nip") and sel is kanan:
            p3 = sel.add_paragraph()
            p3.alignment = ALIGN.CENTER
            p3.add_run("NIP. " + identitas["nip"])


def buat_docx(mentah, identitas=None, template=None, tanggal=None):
    """Pintu masuk tunggal: dict mentah (keluaran AI / Document Schema) -> bytes .docx.
    Memvalidasi + membersihkan dulu; raise SchemaError kalau dokumen tidak layak,
    RendererTidakTersedia kalau python-docx belum terpasang."""
    return render_docx(validasi_dokumen(mentah), identitas=identitas, template=template, tanggal=tanggal)
