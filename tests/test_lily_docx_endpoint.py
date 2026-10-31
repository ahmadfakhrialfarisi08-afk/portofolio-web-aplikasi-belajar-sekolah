"""Test endpoint Lily DOCX lewat APLIKASI UTAMA (app.py), bukan Flask test app buatan.
Tanpa jaringan, tanpa OPENAI_API_KEY: LLM dipalsukan lewat api.lily_docx_routes._pilih_llm, dan urlopen diblokir.
Jalankan: python -m unittest test.test_lily_docx_endpoint -v

Catatan: bila modul lain milik app.py tidak ada di lingkungan test (mis. di CI tanpa stt_engine), modul itu diganti
stub kosong HANYA untuk modul yang memang tidak bisa diimpor. Di project asli semua modul ada, jadi tidak ada yang di-stub.
"""
import importlib
import importlib.util
import io
import os
import shutil
import sys
import tempfile
import types
import unittest
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

os.environ.setdefault('SECRET_KEY', 'kunci-test-saja')
os.environ.pop('OPENAI_API_KEY', None)
_TMP = tempfile.mkdtemp(prefix='lily_docx_test_')
os.environ['LILY_DOCX_DIR'] = os.path.join(_TMP, 'docx')


def _ada(nama):
    try:
        return importlib.util.find_spec(nama) is not None
    except (ImportError, ValueError):
        return False


def _stub_modul_yang_tidak_ada():
    from flask import Blueprint
    if not _ada('image_compress'):
        m = types.ModuleType('image_compress')
        m.GambarTidakValid = type('GambarTidakValid', (Exception,), {})
        m.kompres_gambar_base64 = lambda *a, **k: None
        sys.modules['image_compress'] = m
    if not _ada('stt_engine'):
        m = types.ModuleType('stt_engine')
        m.transkrip_audio_file = lambda *a, **k: ''
        sys.modules['stt_engine'] = m
    if not _ada('api.gemini_routes'):
        m = types.ModuleType('api.gemini_routes')
        m.gemini_bp = Blueprint('gemini_stub', __name__)
        for n in ('set_quiz_context_provider', 'set_maintenance_context_provider', 'set_admin_checker_provider'):
            setattr(m, n, lambda *a, **k: None)
        sys.modules['api.gemini_routes'] = m
    for mod, bp in (('api.scorus_history_routes', 'scorus_history_bp'), ('api.lily_foto_routes', 'lily_foto_bp')):
        if not _ada(mod):
            m = types.ModuleType(mod)
            setattr(m, bp, Blueprint(bp + '_stub', __name__))
            sys.modules[mod] = m


_stub_modul_yang_tidak_ada()
import app as app_utama  # noqa: E402  -- aplikasi Flask yang SEBENARNYA
from api import lily_docx_routes as rute  # noqa: E402
from api import scorus_ai_routes as scorus  # noqa: E402

ISI = "Peserta didik mampu menjelaskan konsep dan menerapkannya dalam praktik secara runtut dan benar."
PILIHAN_A = {"id": "A", "judul": "Lengkap", "deskripsi": "Pendekatan lengkap",
             "bagian": ["Tujuan Pembelajaran", "Kegiatan Pembelajaran", "Asesmen"]}


class LLMPalsu:
    def __init__(self, *jawaban):
        self.jawaban, self.pesan, self.i = list(jawaban), [], 0

    def __call__(self, pesan, maks_token, temperature):
        self.pesan.append(pesan)
        j = self.jawaban[min(self.i, len(self.jawaban) - 1)]
        self.i += 1
        return j if isinstance(j, tuple) else (j, None, {"prompt_tokens": 5, "completion_tokens": 5, "total_tokens": 10})

    @property
    def teks(self):
        return "\n".join(str(m["content"]) for p in self.pesan for m in p)


def analisis_ok():
    def opsi(judul, *bagian):
        return {"judul": judul, "deskripsi": "Pendekatan " + judul, "bagian": list(bagian)}
    return {"dokumen_guru": True, "jenis": "modul_ajar", "topik": "Basis data", "ringkasan": "Modul ajar", "data_kurang": [],
            "pilihan": [opsi("Lengkap", "Tujuan Pembelajaran", "Kegiatan Pembelajaran", "Asesmen"),
                        opsi("Ringkas", "Tujuan", "Proyek", "Rubrik Penilaian")]}


def dokumen_ai():
    return {"jenis": "modul_ajar", "judul": "Modul Ajar Basis Data", "meta": [["Kelas", "XI TKJ"]],
            "bagian": [{"judul": j, "blok": [{"tipe": "paragraf", "teks": ISI + " " + ISI}]} for j in PILIHAN_A["bagian"]],
            "isian_kosong": ["Nomor SK"], "tanda_tangan": True}


class DasarEndpoint(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(dir=_TMP)
        self.dir_docx = os.path.join(self.tmp, 'docx')
        patches = [
            mock.patch.dict(os.environ, {'LILY_DOCX_DIR': self.dir_docx, 'LILY_DOCX_AI': '1'}),
            mock.patch.object(scorus, 'STATUS_FILE', os.path.join(self.tmp, 'status.json')),
            mock.patch.object(scorus, 'USAGE_DIR', os.path.join(self.tmp, 'usage')),
            mock.patch.dict(scorus._status_mem, {'mtime': None, 'data': {}}),
            # jaringan keluar dilarang: kalau ada yang mencoba memanggil OpenAI sungguhan, test langsung gagal
            mock.patch('urllib.request.urlopen', side_effect=AssertionError('panggilan jaringan tidak boleh terjadi di test')),
        ]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)
        rute._hits.clear()
        app_utama.app.config['TESTING'] = True
        self.klien = app_utama.app.test_client()

    def llm(self, *jawaban):
        palsu = LLMPalsu(*jawaban)
        p = mock.patch.object(rute, '_pilih_llm', return_value=palsu)
        p.start()
        self.addCleanup(p.stop)
        return palsu

    def login(self, klien=None, role='guru', username='guru1', **lain):
        klien = klien or self.klien
        with klien.session_transaction() as s:
            s['user'] = dict({'nama': 'Budi Santoso', 'username': username, 'role': role}, **lain)
        return klien

    def post(self, url, data, klien=None):
        return (klien or self.klien).post(url, json=data)


class TestTerhubungKeAppUtama(DasarEndpoint):
    def test_blueprint_terdaftar_sekali_di_app_utama(self):
        aturan = {r.rule: sorted(r.methods - {'HEAD', 'OPTIONS'}) for r in app_utama.app.url_map.iter_rules()
                  if r.rule.startswith('/api/lily/docx/')}
        self.assertEqual(aturan, {'/api/lily/docx/analisis': ['POST'], '/api/lily/docx/buat': ['POST'],
                                  '/api/lily/docx/unduh/<file_id>': ['GET']})
        self.assertIs(app_utama.app.blueprints['lily_docx'], rute.lily_docx_bp)
        with open(os.path.join(ROOT, 'app.py'), encoding='utf-8') as f:
            self.assertEqual(f.read().count('from api.lily_docx_routes import lily_docx_bp'), 1)

    def test_python_docx_ada_di_requirements(self):
        with open(os.path.join(ROOT, 'requirements.txt'), 'rb') as f:
            mentah = f.read()
        teks = mentah.decode('utf-16' if mentah[:2] in (b'\xff\xfe', b'\xfe\xff') else 'utf-8')
        nama = [b.split('==')[0].strip().lower() for b in teks.splitlines() if b.strip()]
        self.assertIn('python-docx', nama)
        self.assertIn('lxml', nama)

    def test_hanya_guru_yang_login(self):
        for url, data in (('/api/lily/docx/analisis', {'prompt': 'modul ajar'}),
                          ('/api/lily/docx/buat', {'prompt': 'x', 'pilihan': PILIHAN_A})):
            r = self.post(url, data)
            self.assertEqual((r.status_code, r.get_json()['kode']), (401, 'belum_login'))
        self.assertEqual(self.klien.get('/api/lily/docx/unduh/' + 'a' * 32).status_code, 401)
        self.login(role='siswa', username='siswa1')
        self.assertEqual(self.post('/api/lily/docx/analisis', {'prompt': 'modul ajar'}).status_code, 401)


class TestAlurLengkap(DasarEndpoint):
    def test_analisis_lalu_buat_lalu_unduh_docx_valid(self):
        self.login()
        llm = self.llm(analisis_ok(), dokumen_ai())
        r = self.post('/api/lily/docx/analisis', {'prompt': 'Buatkan modul ajar basis data kelas XI TKJ'})
        j = r.get_json()
        self.assertEqual((r.status_code, j['success'], j['tahap']), (200, True, 'pilih'))
        self.assertEqual([p['id'] for p in j['pilihan']], ['A', 'B', 'custom'])

        r = self.post('/api/lily/docx/buat', {'prompt': 'Buatkan modul ajar basis data kelas XI TKJ',
                                              'intent': j['intent'], 'pilihan': j['pilihan'][0]})
        b = r.get_json()
        self.assertEqual((r.status_code, b['success']), (200, True), b)
        self.assertEqual(b['nama_file'], 'Modul_Ajar_Basis_Data.docx')
        self.assertEqual(b['peringatan'], [])
        self.assertEqual(llm.i, 2)

        u = self.klien.get(b['url'])
        self.addCleanup(u.close)
        self.assertEqual(u.status_code, 200)
        self.assertEqual(u.mimetype, rute.MIME_DOCX)
        self.assertEqual(u.headers['Cache-Control'], 'no-store')
        self.assertIn('Modul_Ajar_Basis_Data.docx', u.headers['Content-Disposition'])
        self.assertTrue(u.data.startswith(b'PK'))
        import docx
        teks = '\n'.join(p.text for p in docx.Document(io.BytesIO(u.data)).paragraphs)
        self.assertIn('Tujuan Pembelajaran'.upper(), teks.upper())

    def test_identitas_dari_session_masuk_dokumen_tapi_tidak_ke_ai(self):
        self.login(nama='Budi Santoso', sekolah='SMK Contoh Rahasia', nip='1980123')
        llm = self.llm(dokumen_ai())
        r = self.post('/api/lily/docx/buat', {'prompt': 'modul ajar basis data', 'intent': {}, 'pilihan': PILIHAN_A})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertNotIn('Budi Santoso', llm.teks)
        self.assertNotIn('SMK Contoh Rahasia', llm.teks)
        import docx
        u = self.klien.get(r.get_json()['url'])
        self.addCleanup(u.close)
        d = docx.Document(io.BytesIO(u.data))
        semua = '\n'.join(p.text for p in d.paragraphs) + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
        self.assertIn('Budi Santoso', semua)

    def test_file_hanya_bisa_diunduh_pemiliknya_dan_id_ngawur_404(self):
        self.login()
        self.llm(dokumen_ai())
        url = self.post('/api/lily/docx/buat', {'prompt': 'modul ajar', 'intent': {}, 'pilihan': PILIHAN_A}).get_json()['url']
        lain = app_utama.app.test_client()
        self.login(lain, username='guru2')
        r = lain.get(url)
        self.assertEqual((r.status_code, r.get_json()['kode']), (404, 'tidak_ditemukan'))
        for ngawur in ('../../app', 'A' * 32, 'g' * 32, '0' * 31):
            self.assertEqual(self.klien.get('/api/lily/docx/unduh/' + ngawur).status_code, 404, ngawur)
        ok = self.klien.get(url)
        self.addCleanup(ok.close)
        self.assertEqual(ok.status_code, 200)

    def test_pemakaian_api_tercatat_di_log_openai_usage(self):
        self.login()
        self.llm(analisis_ok())
        self.post('/api/lily/docx/analisis', {'prompt': 'modul ajar basis data'})
        berkas = os.listdir(scorus.USAGE_DIR)
        self.assertEqual(len(berkas), 1)
        with open(os.path.join(scorus.USAGE_DIR, berkas[0]), encoding='utf-8') as f:
            baris = f.read()
        self.assertIn('"endpoint": "lily_docx"', baris)
        self.assertIn('"token_total": 10', baris)


class TestPenolakan(DasarEndpoint):
    def test_validasi_input_tidak_memanggil_ai(self):
        self.login()
        llm = self.llm(analisis_ok())
        kasus = [('/api/lily/docx/analisis', None, 400, 'data_kurang'),
                 ('/api/lily/docx/analisis', {'prompt': 123}, 400, 'data_kurang'),
                 ('/api/lily/docx/analisis', {'prompt': ' a '}, 400, 'prompt_kosong'),
                 ('/api/lily/docx/buat', {'prompt': 'modul ajar'}, 400, 'data_kurang'),
                 ('/api/lily/docx/buat', {'prompt': 'modul ajar', 'pilihan': {'id': 'custom'}, 'struktur_custom': ' '}, 400, 'custom_kosong'),
                 ('/api/lily/docx/buat', {'prompt': 'modul ajar', 'pilihan': PILIHAN_A, 'jawaban': 5}, 400, 'data_kurang')]
        for url, data, status, kode in kasus:
            r = self.klien.post(url, json=data) if data is not None else self.klien.post(url, data='bukan json', content_type='text/plain')
            self.assertEqual((r.status_code, r.get_json()['kode']), (status, kode), (url, data))
        self.assertEqual(llm.i, 0)

    def test_pilihan_rusak_dari_browser_ditolak_400(self):
        self.login()
        self.llm(dokumen_ai())
        r = self.post('/api/lily/docx/buat', {'prompt': 'modul ajar', 'intent': {},
                                              'pilihan': {'id': 'A', 'judul': 'X', 'bagian': ['satu']}})
        self.assertEqual((r.status_code, r.get_json()['kode']), (400, 'pilihan_tidak_valid'))

    def test_api_habis_503_tanpa_memanggil_ai(self):
        self.login()
        llm = self.llm(analisis_ok())
        scorus._tandai_api_habis('quota')
        for url, data in (('/api/lily/docx/analisis', {'prompt': 'modul ajar'}),
                          ('/api/lily/docx/buat', {'prompt': 'modul ajar', 'pilihan': PILIHAN_A})):
            r = self.post(url, data)
            j = r.get_json()
            self.assertEqual((r.status_code, j['message_type'], j['api_status']), (503, 'api_exhausted', 'exhausted'))
        self.assertEqual(llm.i, 0)

    def test_quota_dari_ai_menandai_api_habis(self):
        self.login()
        self.llm((None, 'quota', None))
        r = self.post('/api/lily/docx/analisis', {'prompt': 'modul ajar'})
        self.assertEqual((r.status_code, r.get_json()['message_type']), (503, 'api_exhausted'))
        self.assertEqual(scorus._api_habis(), 'quota')

    def test_dimatikan_lewat_env_dan_rate_limit(self):
        self.login()
        llm = self.llm(analisis_ok())
        with mock.patch.dict(os.environ, {'LILY_DOCX_AI': '0'}):
            r = self.post('/api/lily/docx/analisis', {'prompt': 'modul ajar'})
            self.assertEqual((r.status_code, r.get_json()['kode']), (503, 'nonaktif'))
        self.assertEqual(llm.i, 0)
        maks = rute.BATAS_REQUEST['analisis'][0]
        kode = [self.post('/api/lily/docx/analisis', {'prompt': 'modul ajar'}).status_code for _ in range(maks + 1)]
        self.assertEqual(kode[:maks], [200] * maks)
        self.assertEqual(kode[maks], 429)


def tearDownModule():
    shutil.rmtree(_TMP, ignore_errors=True)


if __name__ == '__main__':
    unittest.main()
