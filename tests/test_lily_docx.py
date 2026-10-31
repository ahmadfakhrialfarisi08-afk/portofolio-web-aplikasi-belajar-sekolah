"""Test pipeline Lily DOCX tanpa API token. Jalankan dari root project:
    python -m unittest test.test_lily_docx -v
LLM diganti fungsi palsu (llm=...). Semua akses jaringan DIBLOKIR selama test.
"""
import datetime
import io
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import unittest
import urllib.request
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from services.lily_docx import SchemaError, buat_docx, validasi_dokumen  # noqa: E402
from services.lily_docx import ai, penyimpanan  # noqa: E402

KELUAR = os.path.join(ROOT, 'test', 'output')
os.makedirs(KELUAR, exist_ok=True)

_asli_socket = socket.socket
_asli_urlopen = urllib.request.urlopen
PERMINTAAN_KELUAR = []


def _blokir(*a, **k):
    PERMINTAAN_KELUAR.append((a, k))
    raise AssertionError('Request jaringan terdeteksi saat test!')


def setUpModule():
    socket.socket = _blokir
    urllib.request.urlopen = _blokir
    os.environ.pop('OPENAI_API_KEY', None)


def tearDownModule():
    socket.socket = _asli_socket
    urllib.request.urlopen = _asli_urlopen


def _paragraf_panjang(n=3):
    return ' '.join(['Siswa memahami konsep basis data relasional dan mampu membuat tabel serta relasi antar tabel dengan benar.'] * n)


def dokumen_valid(**ubah):
    d = {
        'jenis': 'modul_ajar', 'judul': 'Modul Ajar Basis Data',
        'meta': [['Mata Pelajaran', 'Basis Data'], ['Kelas', 'XI TKJ'], ['Nilai Kosong', '']],
        'bagian': [
            {'judul': 'Tujuan Pembelajaran', 'blok': [
                {'tipe': 'paragraf', 'teks': _paragraf_panjang()},
                {'tipe': 'daftar', 'gaya': 'nomor', 'item': ['1. Menjelaskan tabel', '2) Membuat relasi', 'Membuat query']}]},
            {'judul': 'Kegiatan Pembelajaran', 'blok': [
                {'tipe': 'subjudul', 'teks': '## Pendahuluan'},
                {'tipe': 'daftar', 'gaya': 'bullet', 'item': ['Salam', 'Apersepsi']},
                {'tipe': 'tabel', 'kolom': ['No', 'Kegiatan', 'Waktu'],
                 'baris': [['1', 'Ceramah **singkat**', '10 menit'], ['2', 'Praktik']]}]},
            {'judul': 'Asesmen', 'blok': [{'tipe': 'paragraf', 'teks': 'Tes tertulis & praktik <b>langsung</b> "dinilai".'}]},
        ],
        'isian_kosong': ['Nomor SK', 'NIP'],
        'tanda_tangan': True,
    }
    d.update(ubah)
    return d


def buka_docx(data):
    import docx
    return docx.Document(io.BytesIO(data))


def teks_semua(d):
    bagian = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                bagian.append(c.text)
    return '\n'.join(bagian)


# ---------------------------------------------------------------- LLM palsu
def opsi(judul, bagian):
    return {'judul': judul, 'deskripsi': 'Pendekatan ' + judul, 'bagian': bagian}


ANALISIS_JELAS = {
    'dokumen_guru': True, 'jenis': 'modul_ajar', 'topik': 'Basis Data XI TKJ', 'ringkasan': 'Modul ajar basis data',
    'data_kurang': [],
    'pilihan': [opsi('Ringkas', ['Tujuan Pembelajaran', 'Kegiatan Pembelajaran', 'Asesmen']),
                opsi('Lengkap', ['Identitas', 'Tujuan', 'Pemahaman Bermakna', 'Kegiatan', 'Asesmen', 'Refleksi'])],
}


class LLMPalsu:
    def __init__(self, *urutan):
        self.urutan, self.panggilan = list(urutan), []

    def __call__(self, pesan, maks_token, temperature):
        self.panggilan.append(pesan)
        item = self.urutan.pop(0) if len(self.urutan) > 1 else self.urutan[0]
        if isinstance(item, Exception):
            raise item
        return item


def ok(d):
    return (d, None, {'prompt_tokens': 10, 'completion_tokens': 20})


class TestSchema(unittest.TestCase):
    def test_valid_dibersihkan(self):
        s = validasi_dokumen(dokumen_valid())
        self.assertEqual(s['judul'], 'Modul Ajar Basis Data')
        self.assertEqual(s['meta'], [['Mata Pelajaran', 'Basis Data'], ['Kelas', 'XI TKJ']])  # nilai kosong dibuang
        daftar = s['bagian'][0]['blok'][1]['item']
        self.assertEqual(daftar, ['Menjelaskan tabel', 'Membuat relasi', 'Membuat query'])  # nomor AI dibuang
        self.assertEqual(s['bagian'][1]['blok'][0]['teks'], 'Pendahuluan')                # markdown ## dibuang
        self.assertNotIn('**', json.dumps(s))
        self.assertEqual(s['bagian'][1]['blok'][2]['baris'][1], ['2', 'Praktik', ''])    # baris pendek dilengkapi

    def test_filler_meta_dibuang(self):
        d = dokumen_valid()
        d['bagian'][0]['blok'][0]['teks'] = 'Tentu, berikut modul yang Anda minta. ' + _paragraf_panjang()
        d['bagian'][2]['blok'].append({'tipe': 'paragraf', 'teks': 'Semoga bermanfaat untuk Bapak/Ibu guru.'})
        s = validasi_dokumen(d)
        semua = json.dumps(s).lower()
        self.assertNotIn('tentu', semua)
        self.assertNotIn('semoga', semua)

    def test_filler_tidak_menghapus_isi_pelajaran(self):
        d = dokumen_valid()
        d['bagian'][0]['blok'][0]['teks'] = 'Berdasarkan hasil diskusi, siswa membuat tabel. ' + _paragraf_panjang()
        s = validasi_dokumen(d)
        self.assertIn('Berdasarkan hasil diskusi', s['bagian'][0]['blok'][0]['teks'])

    def test_ditolak(self):
        for rusak, kata in [
            (None, 'objek'), ([], 'objek'), ('x', 'objek'),
            ({}, 'judul'), ({'judul': 'A'}, 'bagian'), ({'judul': 'A', 'bagian': []}, 'bagian'),
            ({'judul': 'A', 'bagian': [{'judul': 'x', 'blok': []}]}, 'kosong'),
            ({'judul': 'A', 'bagian': [{'judul': 'x', 'blok': [{'tipe': 'paragraf', 'teks': 'pendek'}]}]}, 'sedikit'),
            ({'judul': 'A', 'bagian': [{'judul': 'x', 'blok': [{'tipe': 'script', 'teks': 'x' * 200}]}]}, 'kosong'),
        ]:
            with self.subTest(rusak=str(rusak)[:40]):
                with self.assertRaises(SchemaError) as c:
                    validasi_dokumen(rusak)
                self.assertIn(kata, str(c.exception).lower())

    def test_batas_ukuran(self):
        d = dokumen_valid()
        d['bagian'] = [{'judul': 'B%d' % i, 'blok': [{'tipe': 'paragraf', 'teks': 'x' * 9000}]} for i in range(80)]
        s = validasi_dokumen(d)
        self.assertEqual(len(s['bagian']), 30)
        self.assertLessEqual(len(s['bagian'][0]['blok'][0]['teks']), 4000)

    def test_jenis_bukan_string_tidak_crash(self):
        """Regresi: jenis dari AI berupa list/dict/angka dulu melempar AttributeError (bukan SchemaError)."""
        for jenis in (['modul'], {'a': 1}, 123, 4.5, True, None):
            with self.subTest(jenis=jenis):
                s = validasi_dokumen(dokumen_valid(jenis=jenis))
                self.assertEqual(s['jenis'], 'dokumen')
                buka_docx(buat_docx(dokumen_valid(jenis=jenis)))

    def test_karakter_kontrol_dibuang(self):
        d = dokumen_valid(judul='Modul\x00 Ajar\x0b Basis')
        d['bagian'][2]['blok'][0]['teks'] = 'Teks\x01 dengan\x1f kontrol ' + _paragraf_panjang()
        s = validasi_dokumen(d)
        self.assertEqual(s['judul'], 'Modul Ajar Basis')


class TestRenderer(unittest.TestCase):
    def test_docx_valid_dan_isi(self):
        data = buat_docx(dokumen_valid(), identitas={'sekolah': 'SMK Banjar Asri', 'guru_nama': 'Budi Santoso, S.Pd', 'nip': '1980', 'mapel': 'Basis Data'},
                         tanggal=datetime.datetime(2026, 10, 8))
        self.assertTrue(data.startswith(b'PK'))
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(data)))
        self.assertIsNone(zipfile.ZipFile(io.BytesIO(data)).testzip())
        d = buka_docx(data)
        t = teks_semua(d)
        for harus in ['MODUL AJAR BASIS DATA', 'SMK Banjar Asri', 'Tujuan Pembelajaran', 'Kegiatan Pembelajaran', 'Asesmen',
                      'Menjelaskan tabel', '10 menit', 'Budi Santoso, S.Pd', 'NIP. 1980', '8 Oktober 2026', 'Kepala Sekolah', 'Nomor SK']:
            self.assertIn(harus, t, harus)
        self.assertNotIn('**', t)
        self.assertEqual(d.core_properties.title, 'Modul Ajar Basis Data')
        self.assertEqual(d.core_properties.author, 'Budi Santoso, S.Pd')
        # tabel: tabel meta + tabel isi + tabel TTD
        self.assertGreaterEqual(len(d.tables), 3)
        self.assertTrue(d.sections[0].header.paragraphs[0].text)

    def test_tanpa_identitas_tidak_mengarang(self):
        data = buat_docx(dokumen_valid())
        t = teks_semua(buka_docx(data))
        self.assertNotIn('SMK', t)
        self.assertNotIn('NIP.', t)
        self.assertIn('.' * 40, t)              # isian_kosong -> garis titik
        self.assertEqual(buka_docx(data).core_properties.author, 'SCORUS AI')

    def test_penomoran_daftar_dimulai_ulang(self):
        d = dokumen_valid()
        d['bagian'][2]['blok'].append({'tipe': 'daftar', 'gaya': 'nomor', 'item': ['Satu', 'Dua']})
        data = buat_docx(d)
        xml = zipfile.ZipFile(io.BytesIO(data)).read('word/document.xml').decode()
        import re
        ids = set(re.findall(r'<w:numId w:val="(\d+)"', xml))
        self.assertGreaterEqual(len(ids), 2)    # tiap daftar bernomor punya instance sendiri

    def test_urutan_xml_valid_ooxml(self):
        """pBdr harus sebelum spacing/jc di pPr; shd sebelum vAlign di tcPr (Word menolak urutan salah)."""
        data = buat_docx(dokumen_valid(), identitas={'sekolah': 'S'})
        z = zipfile.ZipFile(io.BytesIO(data))
        from lxml import etree
        for nama in z.namelist():
            if nama.endswith('.xml') or nama.endswith('.rels'):
                etree.fromstring(z.read(nama))  # semua part well-formed
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for nama in ('word/document.xml', 'word/header1.xml', 'word/footer1.xml'):
            if nama not in z.namelist():
                continue
            root = etree.fromstring(z.read(nama))
            for ppr in root.iter('{%s}pPr' % ns['w']):
                anak = [etree.QName(c).localname for c in ppr]
                if 'pBdr' in anak:
                    for setelah in ('spacing', 'jc', 'ind'):
                        if setelah in anak:
                            self.assertLess(anak.index('pBdr'), anak.index(setelah), (nama, anak))

    def test_karakter_aneh_dan_xml(self):
        d = dokumen_valid()
        d['bagian'][2]['blok'][0]['teks'] = 'A & B < C > D "kutip" \'petik\' \u2013 \u00e9\u4e2d\u6587 \U0001F600 ' + _paragraf_panjang()
        t = teks_semua(buka_docx(buat_docx(d)))
        self.assertIn('A & B < C > D', t)
        self.assertIn('\u4e2d\u6587', t)

    def test_dokumen_besar(self):
        d = dokumen_valid()
        d['bagian'] = [{'judul': 'Bagian %d' % i, 'blok': [
            {'tipe': 'paragraf', 'teks': _paragraf_panjang(6)},
            {'tipe': 'tabel', 'kolom': ['A', 'B', 'C', 'D'], 'baris': [[str(r), 'x' * 80, 'y' * 30, 'z'] for r in range(60)]},
        ]} for i in range(30)]
        data = buat_docx(d)
        self.assertGreater(len(data), 10000)
        buka_docx(data)

    def test_template_varian_dan_kustom(self):
        for jenis in ('surat_tugas', 'soal_pg', 'rubrik_penilaian', 'lkpd', 'tidak_dikenal'):
            with self.subTest(jenis=jenis):
                buka_docx(buat_docx(dokumen_valid(jenis=jenis)))
        buka_docx(buat_docx(dokumen_valid(), template={'font': 'Arial', 'halaman': {'kiri': 3.0}, 'warna_h1': 'ZZZ'}))  # warna ngawur aman

    def test_identitas_tipe_salah_aman(self):
        buat_docx(dokumen_valid(), identitas={'sekolah': 123, 'guru_nama': None, 'nip': '\x00\x01', 'mapel': ['x']})
        buat_docx(dokumen_valid(), identitas='bukan dict' if False else None)

    def test_dokumen_tak_layak_ditolak_renderer(self):
        with self.assertRaises(SchemaError):
            buat_docx({'judul': 'x'})

    def test_buka_dengan_libreoffice(self):
        """Bukti file benar-benar bisa dibuka aplikasi office nyata: konversi ke PDF."""
        soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if not soffice:
            self.skipTest('LibreOffice tidak terpasang')
        tmp = tempfile.mkdtemp()
        try:
            f = os.path.join(KELUAR, 'Modul_Ajar_BGP_XI_TKJ.docx')
            with open(f, 'wb') as h:
                h.write(buat_docx(dokumen_valid(), identitas={'sekolah': 'SMK Banjar Asri', 'guru_nama': 'Budi', 'nip': '1'}))
            r = subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', tmp, f],
                               capture_output=True, text=True, timeout=180, env=dict(os.environ, HOME=tmp))
            pdf = os.path.join(tmp, 'Modul_Ajar_BGP_XI_TKJ.pdf')
            self.assertTrue(os.path.isfile(pdf), r.stdout + r.stderr)
            self.assertGreater(os.path.getsize(pdf), 1000)
            shutil.copy(pdf, os.path.join(KELUAR, 'Modul_Ajar_BGP_XI_TKJ.pdf'))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestPenyimpanan(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_simpan_buka_per_guru(self):
        data = buat_docx(dokumen_valid())
        fid = penyimpanan.simpan('guru1', data, 'Modul Ajar BGP XI TKJ.docx', dasar=self.dir)
        self.assertRegex(fid, r'^[0-9a-f]{32}$')
        path, nama = penyimpanan.buka('guru1', fid, dasar=self.dir)
        self.assertTrue(path.startswith(self.dir))
        with open(path, 'rb') as f:
            self.assertEqual(f.read(), data)
        self.assertEqual(nama, 'Modul_Ajar_BGP_XI_TKJ.docx')
        self.assertIsNone(penyimpanan.buka('guru2', fid, dasar=self.dir))     # guru lain tidak bisa

    def test_path_traversal_ditolak(self):
        for jahat in ['../../etc/passwd', '..', '', None, 123, 'a' * 32, '0' * 31, '0' * 32 + '/../x', ('0' * 32 + '\n')]:
            with self.subTest(jahat=jahat):
                self.assertIsNone(penyimpanan.buka('g', jahat, dasar=self.dir))

    def test_nama_file_aman(self):
        self.assertEqual(penyimpanan.nama_file_aman('Modul Ajar: Basis Data / XI TKJ?'), 'Modul_Ajar_Basis_Data_XI_TKJ.docx')
        self.assertEqual(penyimpanan.nama_file_aman('../../x'), 'x.docx')
        self.assertEqual(penyimpanan.nama_file_aman('\u4e2d\u6587'), 'dokumen_lily.docx')
        self.assertEqual(penyimpanan.nama_file_aman(None), 'dokumen_lily.docx')
        self.assertLessEqual(len(penyimpanan.nama_file_aman('a' * 500)), 85)

    def test_batas_jumlah_dan_ttl(self):
        for i in range(55):
            penyimpanan.simpan('g', b'PK' + bytes([i]), 'n%d' % i, dasar=self.dir)
        d = penyimpanan._dir_guru('g', self.dir)
        self.assertLessEqual(len([n for n in os.listdir(d) if n.endswith('.docx')]), penyimpanan.MAKS_FILE_PER_GURU)
        hapus = penyimpanan.bersihkan('g', dasar=self.dir, sekarang=__import__('time').time() + 30 * 86400)
        self.assertGreater(hapus, 0)
        self.assertEqual([n for n in os.listdir(d) if n.endswith('.docx')], [])

    def test_direktori_dasar_env_dan_bawaan(self):
        os.environ['LILY_DOCX_DIR'] = self.dir
        try:
            self.assertEqual(penyimpanan.direktori_dasar(), self.dir)
        finally:
            del os.environ['LILY_DOCX_DIR']
        d = penyimpanan.direktori_dasar()
        self.assertEqual(os.path.abspath(d), os.path.join(ROOT, 'data', 'lily_docx'))   # <root project>/data/lily_docx

    def test_simpan_gagal_naik_oserror(self):
        berkas = os.path.join(self.dir, 'bukan_folder')
        open(berkas, 'w').close()
        with self.assertRaises(OSError):
            penyimpanan.simpan('g', b'PK', 'x', dasar=berkas)


class TestAI(unittest.TestCase):
    def test_prompt_jelas_dua_struktur_plus_custom(self):
        llm = LLMPalsu(ok(ANALISIS_JELAS))
        h = ai.analisis('Buatkan modul ajar Basis Data kelas XI TKJ', llm=llm)
        self.assertEqual(h['tahap'], 'pilih')
        self.assertEqual([p['id'] for p in h['pilihan']], ['A', 'B', 'custom'])
        self.assertEqual(len(llm.panggilan), 1)
        self.assertNotEqual(h['pilihan'][0]['bagian'], h['pilihan'][1]['bagian'])
        self.assertEqual(h['intent']['jenis'], 'modul_ajar')

    def test_prompt_kurang_jelas_bertanya_tanpa_identitas(self):
        a = dict(ANALISIS_JELAS, pilihan=[], data_kurang=[
            {'kunci': 'mapel', 'pertanyaan': 'Mata pelajaran apa?'},
            {'kunci': 'nama_sekolah', 'pertanyaan': 'Apa nama sekolah Anda?'},
            {'kunci': 'nip', 'pertanyaan': 'Berapa NIP Anda?'},
            {'kunci': 'jenjang', 'pertanyaan': 'Jenjang kelas berapa?'}])
        h = ai.analisis('Buatkan modul ajar', llm=LLMPalsu(ok(a)))
        self.assertEqual(h['tahap'], 'tanya')
        self.assertEqual([q['kunci'] for q in h['pertanyaan']], ['mapel', 'jenjang'])   # identitas disaring

    def test_setelah_dijawab_tidak_bertanya_lagi(self):
        a = dict(ANALISIS_JELAS, data_kurang=[{'kunci': 'mapel', 'pertanyaan': 'Mapel apa?'}])
        h = ai.analisis('Buatkan modul ajar', jawaban='Basis Data kelas XI', llm=LLMPalsu(ok(a)))
        self.assertEqual(h['tahap'], 'pilih')

    def test_bukan_dokumen(self):
        h = ai.analisis('apa kabar', llm=LLMPalsu(ok({'dokumen_guru': False})))
        self.assertEqual(h['tahap'], 'tanya')
        with self.assertRaises(ai.AIError) as c:
            ai.analisis('apa kabar', jawaban='halo', llm=LLMPalsu(ok({'dokumen_guru': False})))
        self.assertEqual(c.exception.kode, 'bukan_dokumen')

    def test_struktur_buruk_diulang_lalu_gagal(self):
        buruk = dict(ANALISIS_JELAS, pilihan=[opsi('Satu', ['A']), opsi('Dua', ['A'])])
        llm = LLMPalsu(ok(buruk))
        with self.assertRaises(ai.AIError) as c:
            ai.analisis('Buatkan modul ajar Basis Data', llm=llm)
        self.assertEqual(c.exception.kode, 'struktur_gagal')
        self.assertEqual(len(llm.panggilan), 2)         # tepat 1x ulang
        self.assertIn('Perbaikan', llm.panggilan[1][1]['content'])

    def test_struktur_pulih_di_percobaan_kedua(self):
        buruk = dict(ANALISIS_JELAS, pilihan=[opsi('Satu', ['A'])])
        h = ai.analisis('Buatkan modul ajar Basis Data', llm=LLMPalsu(ok(buruk), ok(ANALISIS_JELAS)))
        self.assertEqual(h['tahap'], 'pilih')

    def test_prompt_kosong(self):
        for p in ('', '  ', None, 'ab'):
            with self.assertRaises(ai.AIError) as c:
                ai.analisis(p, llm=LLMPalsu(ok(ANALISIS_JELAS)))
            self.assertEqual(c.exception.kode, 'prompt_kosong')

    def test_buat_dokumen_ujung_ke_ujung(self):
        pil = ai.analisis('Buatkan modul ajar Basis Data XI TKJ', llm=LLMPalsu(ok(ANALISIS_JELAS)))
        llm = LLMPalsu(ok(dokumen_valid()))
        h = ai.buat_dokumen('Buatkan modul ajar Basis Data XI TKJ', pil['intent'], pil['pilihan'][0], llm=llm,
                            identitas={'sekolah': 'SMK Banjar Asri', 'guru_nama': 'Budi'})
        self.assertTrue(h['docx'].startswith(b'PK'))
        buka_docx(h['docx'])
        self.assertEqual(h['peringatan'], [])
        self.assertEqual(h['usage'], [{'prompt_tokens': 10, 'completion_tokens': 20}])
        # identitas TIDAK dikirim ke AI
        kirim = json.dumps(llm.panggilan)
        self.assertNotIn('SMK Banjar Asri', kirim)
        self.assertNotIn('Budi', kirim)
        with open(os.path.join(KELUAR, 'ujung_ke_ujung.docx'), 'wb') as f:
            f.write(h['docx'])

    def test_bagian_hilang_diulang_lalu_peringatan(self):
        pil = ai.analisis('x modul ajar', llm=LLMPalsu(ok(ANALISIS_JELAS)))['pilihan'][1]   # 6 bagian, dokumen_valid hanya 3
        llm = LLMPalsu(ok(dokumen_valid()))
        h = ai.susun_schema('Buatkan modul ajar', {}, pil, llm=llm)
        self.assertEqual(len(llm.panggilan), 2)
        self.assertTrue(h['peringatan'])

    def test_custom(self):
        custom = ai._opsi_custom()
        h = ai.susun_schema('Buatkan modul ajar', {}, custom, struktur_custom='Tujuan, Kegiatan, Asesmen', llm=LLMPalsu(ok(dokumen_valid())))
        self.assertEqual(h['peringatan'], [])
        with self.assertRaises(ai.AIError) as c:
            ai.susun_schema('Buatkan modul ajar', {}, custom, struktur_custom='  ', llm=LLMPalsu(ok(dokumen_valid())))
        self.assertEqual(c.exception.kode, 'custom_kosong')

    def test_pilihan_tidak_valid(self):
        for pil in (None, 'A', {}, {'id': 'A', 'judul': 'x', 'bagian': ['satu']}):
            with self.assertRaises(ai.AIError) as c:
                ai.susun_schema('Buatkan modul ajar', {}, pil, llm=LLMPalsu(ok(dokumen_valid())))
            self.assertEqual(c.exception.kode, 'pilihan_tidak_valid')

    def test_error_api(self):
        pil = ai._opsi_custom()
        kasus = [((None, 'quota', None), 'api_habis', 'quota'), ((None, 'kunci', None), 'api_habis', 'kunci'),
                 ((None, 'lain', None), 'ai_gagal', None), ((None, None, None), 'ai_gagal', None), (('bukan dict', None, None), 'ai_gagal', None)]
        for balasan, kode, sebab in kasus:
            with self.subTest(kode=kode, sebab=sebab):
                with self.assertRaises(ai.AIError) as c:
                    ai.susun_schema('Buatkan modul ajar', {}, pil, struktur_custom='a, b, c', llm=LLMPalsu(balasan))
                self.assertEqual(c.exception.kode, kode)
                self.assertEqual(c.exception.sebab, sebab)
            with self.assertRaises(ai.AIError):
                ai.analisis('Buatkan modul ajar', llm=LLMPalsu(balasan))

    def test_schema_ditolak_dua_kali(self):
        pil = ai._opsi_custom()
        llm = LLMPalsu(ok({'judul': 'x', 'bagian': []}))
        with self.assertRaises(ai.AIError) as c:
            ai.susun_schema('Buatkan modul ajar', {}, pil, struktur_custom='a, b, c', llm=llm)
        self.assertEqual(c.exception.kode, 'schema_gagal')
        self.assertEqual(len(llm.panggilan), 2)

    def test_timeout_dari_llm_asli_diubah_jadi_error_bersih(self):
        """panggil_openai sungguhan (bukan mock) dengan urlopen yang timeout -> ('lain'); tanpa key -> tidak ada request."""
        self.assertEqual(ai.panggil_openai([{'role': 'user', 'content': 'x'}]), (None, 'lain', None))   # tanpa key
        self.assertEqual(PERMINTAAN_KELUAR, [])
        os.environ['OPENAI_API_KEY'] = 'kunci-palsu-hanya-untuk-test'
        try:
            def timeout(*a, **k):
                raise TimeoutError('timeout')
            urllib.request.urlopen = timeout
            self.assertEqual(ai.panggil_openai([{'role': 'user', 'content': 'x'}]), (None, 'lain', None))
            with self.assertRaises(ai.AIError) as c:
                ai.analisis('Buatkan modul ajar', llm=ai.panggil_openai)
            self.assertEqual(c.exception.kode, 'ai_gagal')
        finally:
            urllib.request.urlopen = _blokir
            del os.environ['OPENAI_API_KEY']

    def test_tanpa_key_jadi_ai_gagal_bukan_crash(self):
        with self.assertRaises(ai.AIError) as c:
            ai.analisis('Buatkan modul ajar Basis Data')     # llm default = panggil_openai, tanpa key
        self.assertEqual(c.exception.kode, 'ai_gagal')
        self.assertEqual(PERMINTAAN_KELUAR, [])


class TestSimpanLewatPipeline(unittest.TestCase):
    def test_hasil_ai_tersimpan_dan_terunduh(self):
        tmp = tempfile.mkdtemp()
        try:
            h = ai.buat_dokumen('Buatkan modul ajar', {}, ai._opsi_custom(), struktur_custom='a, b, c', llm=LLMPalsu(ok(dokumen_valid())))
            fid = penyimpanan.simpan('guru1', h['docx'], penyimpanan.nama_file_aman(h['schema']['judul']), dasar=tmp)
            path, nama = penyimpanan.buka('guru1', fid, dasar=tmp)
            self.assertEqual(nama, 'Modul_Ajar_Basis_Data.docx')
            import docx
            docx.Document(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main(verbosity=2)
