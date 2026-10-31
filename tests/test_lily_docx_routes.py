"""Test integrasi Lily DOCX: prompt -> analisis -> pilihan -> buat -> .docx -> unduh. Tanpa jaringan (LLM dipalsukan),
tanpa menyentuh status API / log usage asli, file ditulis ke folder sementara.
Jalankan: python -m unittest test.test_lily_docx_routes -v"""
import io
import os
import shutil
import sys
import tempfile
import time
import unittest
import zipfile
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flask import Flask

from api import lily_docx_routes as rt
from services.lily_docx import penyimpanan

ISI = "Peserta didik mampu menjelaskan konsep dan menerapkannya dalam praktik secara runtut dan benar."
BAGIAN_A = ["Tujuan Pembelajaran", "Kegiatan Pembelajaran", "Asesmen", "Refleksi"]


def opsi(judul, *bagian):
    return {"judul": judul, "deskripsi": "Pendekatan " + judul, "bagian": list(bagian)}


def analisis_ok(**ubah):
    d = {"dokumen_guru": True, "jenis": "modul_ajar", "topik": "Basis data relasional", "ringkasan": "Modul ajar basis data",
         "data_kurang": [], "pilihan": [opsi("Lengkap", *BAGIAN_A), opsi("Ringkas Proyek", "Tujuan", "Proyek", "Rubrik Penilaian")]}
    d.update(ubah)
    return d


def dokumen_ai(*judul_bagian, **ubah):
    d = {"jenis": "modul_ajar", "judul": "Modul Ajar Basis Data: Kelas XI", "meta": [["Kelas", "XI TKJ"]],
         "bagian": [{"judul": j, "blok": [{"tipe": "paragraf", "teks": ISI + " " + ISI}]} for j in judul_bagian],
         "isian_kosong": ["Nama guru", "NIP", "Nomor SK"], "tanda_tangan": True}
    d.update(ubah)
    return d


class LLMPalsu:
    def __init__(self, *jawaban):
        self.jawaban, self.pesan, self.i = list(jawaban), [], 0

    def __call__(self, pesan, maks_token, temperature):
        self.pesan.append(pesan)
        j = self.jawaban[min(self.i, len(self.jawaban) - 1)]
        self.i += 1
        return j if isinstance(j, tuple) else (j, None, {"total_tokens": 10})

    @property
    def teks(self):
        return "\n".join(str(m["content"]) for p in self.pesan for m in p)


def teks_dokumen(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    return d, "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])


class Dasar(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="lily_docx_uji_")
        self.env = mock.patch.dict(os.environ, {"LILY_DOCX_DIR": self.dir, "LILY_DOCX_AI": "1"})
        self.env.start()
        self.p = [mock.patch.object(rt, "_api_habis", return_value=None),
                  mock.patch.object(rt, "_tandai_api_habis"), mock.patch.object(rt, "_tandai_api_pulih"),
                  mock.patch.object(rt, "_catat_usage")]
        self.m = [p.start() for p in self.p]
        rt._hits.clear()
        app = Flask(__name__)
        app.secret_key = "uji"
        app.register_blueprint(rt.lily_docx_bp)
        self.app, self.c = app, app.test_client()
        self.login()

    def tearDown(self):
        for p in self.p:
            p.stop()
        self.env.stop()
        shutil.rmtree(self.dir, ignore_errors=True)

    def login(self, user=None, client=None):
        with (client or self.c).session_transaction() as s:
            s["user"] = user if user is not None else {"role": "guru", "username": "guru1"}

    def pakai(self, *jawaban):
        llm = LLMPalsu(*jawaban)
        pat = mock.patch.object(rt, "_pilih_llm", return_value=llm)
        pat.start()
        self.addCleanup(pat.stop)
        return llm

    def post(self, path, body=None, client=None):
        return (client or self.c).post(path, json=body)

    def unduh(self, url, client=None):
        r = (client or self.c).get(url)
        data = r.data
        r.close()
        return data

    def berkas_tersimpan(self):
        return [f for _, _, fs in os.walk(self.dir) for f in fs]


class TestAkses(Dasar):
    def test_wajib_login_guru(self):
        llm = self.pakai(analisis_ok())
        for user in (None, {"role": "siswa", "username": "s1"}, {"role": "guru"}):
            c = self.app.test_client()
            if user:
                self.login(user, c)
            for r in (c.post("/api/lily/docx/analisis", json={"prompt": "modul ajar"}),
                      c.post("/api/lily/docx/buat", json={}), c.get("/api/lily/docx/unduh/" + "a" * 32)):
                self.assertEqual(r.status_code, 401)
                self.assertEqual(r.get_json()["kode"], "belum_login")
        self.assertEqual(llm.i, 0)

    def test_dimatikan_lewat_env(self):
        llm = self.pakai(analisis_ok())
        with mock.patch.dict(os.environ, {"LILY_DOCX_AI": "0"}):
            r = self.post("/api/lily/docx/analisis", {"prompt": "modul ajar"})
        self.assertEqual((r.status_code, r.get_json()["kode"]), (503, "nonaktif"))
        self.assertEqual(llm.i, 0)

    def test_hanya_rute_docx_yang_ditambahkan(self):
        aturan = {r.rule for r in self.app.url_map.iter_rules() if r.endpoint != "static"}
        self.assertTrue(all(a.startswith("/api/lily/docx/") for a in aturan), aturan)


class TestAlur(Dasar):
    def test_prompt_pilihan_generate_unduh(self):
        llm = self.pakai(analisis_ok(), dokumen_ai(*BAGIAN_A))
        prompt = "Buatkan modul ajar basis data kelas XI TKJ"
        a = self.post("/api/lily/docx/analisis", {"prompt": prompt}).get_json()
        self.assertTrue(a["success"])
        self.assertEqual(a["tahap"], "pilih")
        self.assertEqual([o["id"] for o in a["pilihan"]], ["A", "B", "custom"])

        r = self.post("/api/lily/docx/buat", {"prompt": prompt, "intent": a["intent"], "pilihan": a["pilihan"][0]})
        b = r.get_json()
        self.assertEqual(r.status_code, 200, b)
        self.assertTrue(b["success"])
        self.assertEqual(b["peringatan"], [])
        self.assertEqual(b["nama_file"], "Modul_Ajar_Basis_Data_Kelas_XI.docx")
        self.assertEqual(b["url"], "/api/lily/docx/unduh/" + b["file_id"])
        self.assertEqual(llm.i, 2)

        u = self.c.get(b["url"])
        self.addCleanup(u.close)
        self.assertEqual(u.status_code, 200)
        self.assertEqual(u.mimetype, rt.MIME_DOCX)
        self.assertIn("Modul_Ajar_Basis_Data_Kelas_XI.docx", u.headers["Content-Disposition"])
        self.assertIn("attachment", u.headers["Content-Disposition"])
        self.assertEqual(u.headers["Cache-Control"], "no-store")
        self.assertTrue(u.data.startswith(b"PK"))
        zipfile.ZipFile(io.BytesIO(u.data)).testzip()
        d, teks = teks_dokumen(u.data)
        self.assertIn("modul ajar basis data", teks.lower())
        for bagian in BAGIAN_A:
            self.assertIn(bagian, teks)

    def test_data_kurang_bertanya_lalu_lanjut_dengan_jawaban(self):
        kurang = analisis_ok(data_kurang=[{"kunci": "mata_pelajaran", "pertanyaan": "Mata pelajarannya apa?"}], pilihan=[])
        llm = self.pakai(kurang, analisis_ok())
        a = self.post("/api/lily/docx/analisis", {"prompt": "buatkan modul ajar"}).get_json()
        self.assertEqual((a["tahap"], a["pertanyaan"][0]["kunci"]), ("tanya", "mata_pelajaran"))
        self.assertNotIn("pilihan", a)
        a2 = self.post("/api/lily/docx/analisis", {"prompt": "buatkan modul ajar", "jawaban": "Basis data, kelas XI"}).get_json()
        self.assertEqual(a2["tahap"], "pilih")
        self.assertIn("Basis data, kelas XI", llm.teks)

    def test_custom_wajib_isi_lalu_berhasil(self):
        llm = self.pakai(dokumen_ai("Latar Belakang", "Langkah Praktik"))
        body = {"prompt": "modul ajar", "intent": {}, "pilihan": {"id": "custom"}}
        r = self.post("/api/lily/docx/buat", body)
        self.assertEqual((r.status_code, r.get_json()["kode"]), (400, "custom_kosong"))
        self.assertEqual(llm.i, 0)
        body["struktur_custom"] = "1. Latar Belakang\n2. Langkah Praktik"
        r = self.post("/api/lily/docx/buat", body)
        self.assertEqual(r.status_code, 200, r.get_json())
        _, teks = teks_dokumen(self.unduh(r.get_json()["url"]))
        self.assertIn("Langkah Praktik", teks)

    def test_identitas_hanya_dari_session_dan_tidak_dikirim_ke_ai(self):
        llm = self.pakai(dokumen_ai(*BAGIAN_A), dokumen_ai(*BAGIAN_A))
        body = {"prompt": "modul ajar", "intent": {}, "pilihan": {"id": "A", **opsi("Lengkap", *BAGIAN_A)}}
        # tanpa identitas di session: tidak ada nama dikarang, penulis dokumen bawaan, isian dikosongkan
        d, teks = teks_dokumen(self.unduh(self.post("/api/lily/docx/buat", body).get_json()["url"]))
        self.assertEqual(d.core_properties.author, "SCORUS AI")
        self.assertNotIn("Budi", teks)
        self.assertIn("Nomor SK", teks)
        # dengan identitas di session: dipakai renderer, tetap tidak dikirim ke AI
        self.login({"role": "guru", "username": "guru1", "nama": "Budi Santoso", "nip": "19800101", "sekolah": "SMK Contoh Uji"})
        d, teks = teks_dokumen(self.unduh(self.post("/api/lily/docx/buat", body).get_json()["url"]))
        self.assertEqual(d.core_properties.author, "Budi Santoso")
        self.assertIn("Budi Santoso", teks)
        for rahasia in ("Budi Santoso", "19800101", "SMK Contoh Uji"):
            self.assertNotIn(rahasia, llm.teks)

    def test_bagian_hilang_jadi_peringatan_bukan_gagal_diam_diam(self):
        self.pakai(dokumen_ai("Tujuan Pembelajaran", "Kegiatan Pembelajaran", "Asesmen"))
        body = {"prompt": "modul ajar", "intent": {}, "pilihan": {"id": "A", **opsi("Lengkap", *BAGIAN_A)}}
        b = self.post("/api/lily/docx/buat", body).get_json()
        self.assertTrue(b["success"])
        self.assertEqual(b["peringatan"], ['Bagian "Refleksi" tidak berhasil dibuat.'])


class TestKegagalan(Dasar):
    BODY = {"prompt": "modul ajar", "intent": {}, "pilihan": {"id": "A", **opsi("Lengkap", *BAGIAN_A)}}

    def test_input_tidak_valid(self):
        self.pakai(analisis_ok())
        for path, body in (("analisis", None), ("analisis", {"prompt": 5}), ("analisis", {"prompt": "x", "jawaban": 3}),
                           ("buat", {"prompt": "modul", "pilihan": "A"}), ("buat", {"prompt": ["x"], "pilihan": {}}),
                           ("buat", {"prompt": "modul", "pilihan": {}, "struktur_custom": 7})):
            r = self.post("/api/lily/docx/" + path, body)
            self.assertEqual((r.status_code, r.get_json()["kode"]), (400, "data_kurang"), (path, body))
        r = self.post("/api/lily/docx/analisis", {"prompt": "  "})
        self.assertEqual(r.get_json()["kode"], "prompt_kosong")
        r = self.c.post("/api/lily/docx/analisis", data="bukan json", content_type="application/json")
        self.assertEqual(r.status_code, 400)

    def test_pilihan_palsu_dari_browser_ditolak(self):
        llm = self.pakai(dokumen_ai(*BAGIAN_A))
        r = self.post("/api/lily/docx/buat", {"prompt": "modul ajar", "pilihan": {"id": "A", "judul": "X", "bagian": ["satu"]}})
        self.assertEqual((r.status_code, r.get_json()["kode"]), (400, "pilihan_tidak_valid"))
        self.assertEqual(llm.i, 0)

    def test_schema_invalid_ditangani_dan_tidak_menyimpan_file(self):
        self.pakai({"judul": "X", "bagian": []})
        r = self.post("/api/lily/docx/buat", self.BODY)
        self.assertEqual((r.status_code, r.get_json()["kode"]), (502, "schema_gagal"))
        self.assertFalse(r.get_json()["success"])
        self.assertEqual(self.berkas_tersimpan(), [])

    def test_ai_tidak_merespons(self):
        self.pakai((None, "lain", None))
        r = self.post("/api/lily/docx/analisis", {"prompt": "modul ajar"})
        self.assertEqual((r.status_code, r.get_json()["kode"]), (502, "ai_gagal"))

    def test_kuota_habis_ditandai_dan_dijawab_jujur(self):
        self.pakai((None, "quota", None))
        with mock.patch.object(rt, "_teks_api_habis", return_value="API habis"):
            r = self.post("/api/lily/docx/analisis", {"prompt": "modul ajar"})
        j = r.get_json()
        self.assertEqual((r.status_code, j["kode"], j["api_status"]), (503, "api_habis", "exhausted"))
        self.m[1].assert_called_once_with("quota")      # _tandai_api_habis

    def test_api_sudah_ditandai_habis_tidak_menghubungi_openai(self):
        llm = self.pakai(analisis_ok())
        with mock.patch.object(rt, "_api_habis", return_value="quota"), mock.patch.object(rt, "_teks_api_habis", return_value="API habis"):
            r = self.post("/api/lily/docx/analisis", {"prompt": "modul ajar"})
        self.assertEqual((r.status_code, r.get_json()["message_type"]), (503, "api_exhausted"))
        self.assertEqual(llm.i, 0)

    def test_renderer_gagal_tidak_menyimpan_file(self):
        self.pakai(dokumen_ai(*BAGIAN_A))
        with mock.patch.object(rt, "buat_docx", side_effect=RuntimeError("rusak")):
            r = self.post("/api/lily/docx/buat", self.BODY)
        self.assertEqual((r.status_code, r.get_json()["kode"]), (500, "render_gagal"))
        with mock.patch.object(rt, "buat_docx", return_value=b"bukan docx"):
            r = self.post("/api/lily/docx/buat", self.BODY)
        self.assertEqual(r.get_json()["kode"], "render_gagal")
        self.assertEqual(self.berkas_tersimpan(), [])

    def test_gagal_simpan(self):
        self.pakai(dokumen_ai(*BAGIAN_A))
        with mock.patch.object(rt.penyimpanan, "simpan", side_effect=OSError("disk penuh")):
            r = self.post("/api/lily/docx/buat", self.BODY)
        self.assertEqual((r.status_code, r.get_json()["kode"]), (500, "simpan_gagal"))

    def test_pencatatan_usage_yang_rusak_tidak_menggagalkan_dokumen(self):
        self.pakai(dokumen_ai(*BAGIAN_A))
        self.m[3].side_effect = RuntimeError("log rusak")
        r = self.post("/api/lily/docx/buat", self.BODY)
        self.assertEqual(r.status_code, 200)

    def test_rate_limit(self):
        self.pakai(analisis_ok())
        with mock.patch.dict(rt.BATAS_REQUEST, {"analisis": (2, 600)}):
            kode = [self.post("/api/lily/docx/analisis", {"prompt": "modul ajar"}).status_code for _ in range(3)]
            self.assertEqual(kode, [200, 200, 429])
            self.assertEqual(self.post("/api/lily/docx/analisis", {"prompt": "x"}).status_code, 400)  # tidak makan jatah


class TestUnduh(Dasar):
    def buat(self, client=None):
        self.pakai(dokumen_ai(*BAGIAN_A))
        body = {"prompt": "modul ajar", "intent": {}, "pilihan": {"id": "A", **opsi("Lengkap", *BAGIAN_A)}}
        return self.post("/api/lily/docx/buat", body, client).get_json()

    def test_guru_lain_tidak_bisa_mengunduh(self):
        b = self.buat()
        lain = self.app.test_client()
        self.login({"role": "guru", "username": "guru2"}, lain)
        self.assertEqual(lain.get(b["url"]).status_code, 404)
        ok = self.c.get(b["url"])
        self.assertEqual(ok.status_code, 200)
        ok.close()

    def test_file_id_aneh_ditolak(self):
        self.buat()
        for fid in ("..%2F..%2Fapp.py", "A" * 32, "abc", "a" * 33, "0" * 32, "%00", "..", "a" * 31 + "/"):
            self.assertEqual(self.c.get("/api/lily/docx/unduh/" + fid).status_code, 404, fid)


class TestPenyimpanan(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="lily_simpan_uji_")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_nama_file_masuk_akal(self):
        f = penyimpanan.nama_file_aman
        self.assertEqual(f("Modul Ajar Basis Data: XI TKJ"), "Modul_Ajar_Basis_Data_XI_TKJ.docx")
        self.assertEqual(f("Surat Tugas / Workshop ../../x"), "Surat_Tugas_Workshop_x.docx")
        self.assertEqual(f("Évaluasi Ünit"), "Evaluasi_Unit.docx")
        self.assertEqual(f("???"), "dokumen_lily.docx")
        self.assertEqual(f(None), "dokumen_lily.docx")
        self.assertLessEqual(len(f("a" * 500)), 85)

    def test_simpan_buka_dan_tanpa_sisa_tmp(self):
        fid = penyimpanan.simpan("guru1", b"PKdata", "Dok.docx", dasar=self.dir)
        path, nama = penyimpanan.buka("guru1", fid, dasar=self.dir)
        with open(path, "rb") as f:
            isi = f.read()
        self.assertEqual((isi, nama), (b"PKdata", "Dok.docx"))
        self.assertIsNone(penyimpanan.buka("guru2", fid, dasar=self.dir))
        sisa = [f for _, _, fs in os.walk(self.dir) for f in fs]
        self.assertFalse([f for f in sisa if f.endswith(".tmp")])
        self.assertNotIn("guru1", "".join(sisa) + "".join(os.listdir(self.dir)))   # folder memakai hash, bukan username

    def test_pembersihan_kedaluwarsa_dan_batas_jumlah(self):
        lama = penyimpanan.simpan("g", b"PK1", "a.docx", dasar=self.dir)
        d = penyimpanan._dir_guru("g", self.dir)
        t = time.time() - 30 * 86400
        for f in os.listdir(d):
            os.utime(os.path.join(d, f), (t, t))
        baru = penyimpanan.simpan("g", b"PK2", "b.docx", dasar=self.dir)      # simpan memicu pembersihan
        self.assertIsNone(penyimpanan.buka("g", lama, dasar=self.dir))
        self.assertIsNotNone(penyimpanan.buka("g", baru, dasar=self.dir))
        ids = [penyimpanan.simpan("g", b"PK", "c.docx", dasar=self.dir) for _ in range(5)]
        penyimpanan.bersihkan("g", dasar=self.dir, maks=3)
        ada = [i for i in ids + [baru] if penyimpanan.buka("g", i, dasar=self.dir)]
        self.assertEqual(len(ada), 2)                                         # maks-1 disisakan untuk file berikutnya
        self.assertIn(ids[-1], ada)


if __name__ == "__main__":
    unittest.main()
